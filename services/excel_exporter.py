"""
Module responsible for exporting and updating lead time and revenue results directly in Excel and CSV files.
Drop-off files receive operational Lead Time (business hours) and financial columns.
Collection files receive financial columns without Lead Time.
"""

from pathlib import Path
from typing import Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import pandas as pd

from config import (
    COL_PACKAGE_TYPE,
    COL_UNIT_REVENUE,
    COL_CUMULATIVE_REVENUE,
    COL_LEAD_TIME_BUSINESS_FORMATTED,
    COL_LEAD_TIME_BUSINESS_MINUTES,
    COL_LEAD_TIME_BUSINESS_HOURS,
    COL_LEAD_TIME_RAW_FORMATTED,
    COL_PROCESSING_STATUS,
    SUMMARY_SHEET_NAME,
)


class ExcelExporter:
    """
    Handles writing lead time calculations, pricing, and summaries to Excel and CSV files.
    """

    def __init__(self, default_output_dir: Path):
        self.default_output_dir = Path(default_output_dir)
        self.default_output_dir.mkdir(parents=True, exist_ok=True)

    def update_dropoff_in_place(
        self,
        file_path: Path,
        enriched_df: pd.DataFrame,
        operational_metrics: Dict[str, Any],
        financial_metrics: Dict[str, Any],
    ) -> Path:
        """
        Updates the original Drop-off Excel file in-place by adding/updating calculated columns
        and refreshing the 'Resumo_Metricas' sheet.
        """
        target_path = Path(file_path).resolve()

        try:
            wb = openpyxl.load_workbook(target_path)
        except PermissionError as exc:
            raise PermissionError(
                f"O arquivo '{target_path.name}' está aberto no Excel ou em outro programa. "
                f"Por favor, feche o arquivo e tente novamente."
            ) from exc

        data_ws = wb.active if "Sheet1" not in wb.sheetnames else wb["Sheet1"]

        target_columns = [
            COL_PACKAGE_TYPE,
            COL_UNIT_REVENUE,
            COL_CUMULATIVE_REVENUE,
            COL_LEAD_TIME_BUSINESS_FORMATTED,
            COL_LEAD_TIME_BUSINESS_MINUTES,
            COL_LEAD_TIME_BUSINESS_HOURS,
            COL_LEAD_TIME_RAW_FORMATTED,
            COL_PROCESSING_STATUS,
        ]

        existing_headers = {}
        for col_idx in range(1, data_ws.max_column + 1):
            val = data_ws.cell(row=1, column=col_idx).value
            if val is not None:
                existing_headers[str(val).strip()] = col_idx

        col_mapping = {}
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill_fin = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid")
        header_fill_ops = PatternFill(start_color="2F4F4F", end_color="2F4F4F", fill_type="solid")

        current_max = data_ws.max_column
        for target_col in target_columns:
            if target_col in existing_headers:
                col_mapping[target_col] = existing_headers[target_col]
            else:
                current_max += 1
                col_mapping[target_col] = current_max
                cell = data_ws.cell(row=1, column=current_max, value=target_col)
                cell.font = header_font
                cell.fill = header_fill_fin if "Remuneração" in target_col or "Tipo" in target_col else header_fill_ops
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_idx, (_, row_data) in enumerate(enriched_df.iterrows(), start=2):
            for target_col in target_columns:
                val = row_data[target_col]
                cell_value = None if pd.isna(val) else val
                target_col_idx = col_mapping[target_col]
                cell = data_ws.cell(row=row_idx, column=target_col_idx, value=cell_value)
                if "Remuneração" in target_col and cell_value is not None:
                    cell.number_format = '"R$" #,##0.00'

        # Refresh Resumo_Metricas
        if SUMMARY_SHEET_NAME in wb.sheetnames:
            del wb[SUMMARY_SHEET_NAME]

        summary_ws = wb.create_sheet(title=SUMMARY_SHEET_NAME)
        self._format_summary_sheet(summary_ws, target_path.name, financial_metrics, operational_metrics)

        try:
            wb.save(target_path)
            wb.close()
        except PermissionError as exc:
            raise PermissionError(
                f"Não foi possível salvar o arquivo '{target_path.name}' porque ele está em uso no Excel."
            ) from exc

        return target_path

    def update_collection_in_place(self, file_path: Path, enriched_df: pd.DataFrame) -> Path:
        """
        Updates the original Collection file (.csv or .xlsx) in-place with financial columns.
        (Lead Time is excluded from Collection orders as it depends on the buyer).
        """
        target_path = Path(file_path).resolve()
        try:
            if target_path.suffix.lower() == ".csv":
                enriched_df.to_csv(target_path, index=False, encoding="utf-8-sig")
            else:
                enriched_df.to_excel(target_path, index=False, engine="openpyxl")
        except PermissionError as exc:
            raise PermissionError(
                f"O arquivo de retiradas '{target_path.name}' está aberto. Por favor, feche-o e tente novamente."
            ) from exc
        return target_path

    def export_monthly_report(
        self,
        reference_month: str,
        operational_metrics: Dict[str, Any],
        financial_metrics: Dict[str, Any],
        df_dropoff: pd.DataFrame,
        df_collection: pd.DataFrame,
    ) -> Path:
        """
        Exports a consolidated multi-sheet Excel workbook for an entire historical billing month.
        """
        out_filename = f"relatorio_mensal_{reference_month.replace('-', '_')}.xlsx"
        target_path = self.default_output_dir / out_filename

        wb = openpyxl.Workbook()
        ws_summary = wb.active
        ws_summary.title = SUMMARY_SHEET_NAME
        self._format_summary_sheet(ws_summary, f"Mês Consolidado: {reference_month}", financial_metrics, operational_metrics)

        # Drop-off sheet
        if not df_dropoff.empty:
            ws_drop = wb.create_sheet(title="Drop_off_Pacotes")
            ws_drop.append(list(df_dropoff.columns))
            for _, r in df_dropoff.iterrows():
                ws_drop.append([None if pd.isna(v) else v for v in r])

        # Collection sheet
        if not df_collection.empty:
            ws_coll = wb.create_sheet(title="Retiradas_Pacotes")
            ws_coll.append(list(df_collection.columns))
            for _, r in df_collection.iterrows():
                ws_coll.append([None if pd.isna(v) else v for v in r])

        wb.save(target_path)
        wb.close()
        return target_path

    def _format_summary_sheet(
        self,
        ws: openpyxl.worksheet.worksheet.Worksheet,
        filename: str,
        financial_metrics: Dict[str, Any],
        operational_metrics: Dict[str, Any],
    ) -> None:
        """Applies styles to summary worksheet."""
        title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
        sub_font = Font(name="Calibri", size=10, italic=True, color="595959")
        section_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        section_fill_fin = PatternFill(start_color="198754", end_color="198754", fill_type="solid")
        section_fill_ops = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        tbl_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        tbl_header_fill = PatternFill(start_color="495057", end_color="495057", fill_type="solid")
        bold_font = Font(name="Calibri", size=11, bold=True)
        regular_font = Font(name="Calibri", size=11)
        highlight_font = Font(name="Calibri", size=12, bold=True, color="198754")
        highlight_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")

        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        ws["A1"] = "Relatório Executivo Shopee - Ponto de Coleta e Retirada (PUDO)"
        ws["A1"].font = title_font
        ws["A2"] = f"Arquivo de Origem: {filename}  |  Horário Comercial: Seg-Sex 08h-19h30, Sáb 09h-15h"
        ws["A2"].font = sub_font

        curr_row = 4

        # Section 1: Financial
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=2)
        c = ws.cell(row=curr_row, column=1, value="ESTIMATIVA DE REMUNERAÇÃO (GANHOS DA AGÊNCIA)")
        c.font = section_font
        c.fill = section_fill_fin
        c.alignment = Alignment(horizontal="center", vertical="center")
        curr_row += 1

        ws.cell(row=curr_row, column=1, value="Indicador Financeiro").font = tbl_header_font
        ws.cell(row=curr_row, column=1).fill = tbl_header_fill
        ws.cell(row=curr_row, column=2, value="Valor").font = tbl_header_font
        ws.cell(row=curr_row, column=2).fill = tbl_header_fill
        ws.cell(row=curr_row, column=2).alignment = Alignment(horizontal="center", vertical="center")
        curr_row += 1

        for label, val in financial_metrics.items():
            c1 = ws.cell(row=curr_row, column=1, value=label)
            c2 = ws.cell(row=curr_row, column=2, value=val)
            c1.border = thin_border
            c2.border = thin_border
            c2.alignment = Alignment(horizontal="center", vertical="center")

            if "FATURAMENTO TOTAL" in label or "Volume Total" in label:
                c1.font = highlight_font
                c2.font = highlight_font
                c1.fill = highlight_fill
                c2.fill = highlight_fill
            else:
                c1.font = bold_font if ("Faturamento" in label or "Ticket" in label or "Total" in label) else regular_font
                c2.font = regular_font
            curr_row += 1

        curr_row += 1

        # Section 2: Operational Lead Time in Business Hours (Drop-off only)
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=2)
        c = ws.cell(
            row=curr_row, column=1, value="LEAD TIME ÚTIL (DROP-OFF: POSTAGENS & DEVOLUÇÕES ATÉ O CAMINHÃO)"
        )
        c.font = section_font
        c.fill = section_fill_ops
        c.alignment = Alignment(horizontal="center", vertical="center")
        curr_row += 1

        ws.cell(row=curr_row, column=1, value="Métrica Operacional").font = tbl_header_font
        ws.cell(row=curr_row, column=1).fill = tbl_header_fill
        ws.cell(row=curr_row, column=2, value="Resultado").font = tbl_header_font
        ws.cell(row=curr_row, column=2).fill = tbl_header_fill
        ws.cell(row=curr_row, column=2).alignment = Alignment(horizontal="center", vertical="center")
        curr_row += 1

        for label, val in operational_metrics.items():
            c1 = ws.cell(row=curr_row, column=1, value=label)
            c2 = ws.cell(row=curr_row, column=2, value=val)
            c1.font = bold_font if "Lead Time Útil Médio" in label else regular_font
            c2.font = regular_font
            c1.border = thin_border
            c2.border = thin_border
            c2.alignment = Alignment(horizontal="center", vertical="center")
            curr_row += 1

        ws.column_dimensions["A"].width = 52
        ws.column_dimensions["B"].width = 28
